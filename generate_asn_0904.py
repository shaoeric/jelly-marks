#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
ASN 上传表生成器（ASN_Template）
================================

把「唛头 Excel」（每个 sheet = 1 个托盘）展开成可上传 ASN 系统的
「ASN_Template」表格：每个托盘的每个 part-number 占一行，列名/顺序固定 A-AE。

输入
----
1. 唛头 Excel：每个 sheet 一个托盘。
   - A8/B8 = LOT NUMBER，格式「托盘号-托盘总数」，托盘号 = 该 sheet 在工作簿中的顺序（第 1 个 sheet 就是第 1 个托盘）
   - 第 9 行是 part 表头，第 10 行起：A 列 part-number、C 列数量
   - 之后依次是 WEIGHT (IN KG)（B 列总重）、DIMENSIONS (MM) LxWxH（B 列尺寸）
2. ASN 导出 Excel 中的数据 sheet：A-AE 列，含每个 part-number 的静态字段。
3. ASN 导出 Excel 中的单位重量 sheet：无表头 2 列，A = part-number，B = 单位重量（正数）。

输出
----
新建 Excel，仅含一个名为 ASN_Template 的 sheet：
- AB package_id              ← 托盘号（按 sheet 顺序 1..N）
- S  allowed_qty             ← 数据 sheet
- V  items_per_package       ← 唛头中该托盘对应 part 的数量
- W  net_weight_per_package  = 数量 × 单位重量（只写在每个托盘的首行，其余行留空）
- X  gross_weight_per_package← 唛头中该托盘的总重（只写在每个托盘的首行，其余行留空）
- Z  dimensions              ← 唛头中该托盘的尺寸（x 统一写成 *）
其余列沿用数据 sheet 的值。

用法
----
    uv run python generate_asn_0904.py \\
        --mark-path "cases/case2/Wichita Batch4-唛头(2).xlsx" \\
        --asn-path "cases/case2/ASN_Bulk_Template_Batch4 9-24_核对修正.xlsx" \\
        --asn-sheet "ASN_Template (2)" \\
        --weight-sheet "Sheet1" \\
        --out "cases/case2/ASN_Batch4_上传.xlsx"

任一校验不通过时会列出全部问题并中止（不生成输出文件）。
"""

import argparse
import os
import re
import sys
from collections import Counter

try:
    import openpyxl
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("missing openpyxl, run: pip install openpyxl")


# ----------------------------------------------------------------------------
# 固定列名（顺序即 A-AE 列顺序，直接写死）
# ----------------------------------------------------------------------------
OUT_SHEET = "ASN_Template"

HEADERS = [
    "asn_import", "asn_status", "project_number", "supplier", "remarks",
    "sfc_number", "cargo_ready_date", "shipment_date", "container_number",
    "container_type", "portal_mode_of_transport", "plate_number",
    "lorry_receipt_number", "tracking_number", "allocation_code",
    "part_number", "allocation_week", "po", "allowed_qty", "package_type",
    "package_quantity", "items_per_package", "net_weight_per_package",
    "gross_weight_per_package", "unit_weight", "dimensions",
    "dimensions_unit", "package_id", "count_of_boxes", "country_of_origin",
    "source",
]
COL = {name: i for i, name in enumerate(HEADERS)}

# 数值列：写入时把数字字符串转成数字、空字符串转成空单元格（与参考表一致）
NUMERIC_COLS = {
    "asn_import", "allowed_qty", "package_quantity", "items_per_package",
    "net_weight_per_package", "gross_weight_per_package", "package_id",
    "count_of_boxes",
}

# 唛头 sheet 内的固定位置/标签
LOT_ROW = 8                     # A8/B8 = LOT NUMBER
FIRST_PART_ROW = 10             # 第 9 行是 part 表头
LOT_LABEL = "LOT NUMBER"
WEIGHT_LABEL = "WEIGHT (IN KG)"
DIMS_LABEL = "DIMENSIONS (MM) LxWxH"

LOT_RE = re.compile(r"^(\d+)\s*-\s*(\d+)$")


# ----------------------------------------------------------------------------
# 小工具
# ----------------------------------------------------------------------------
def text(value):
    """单元格值转成去空白的字符串（None -> ""）。"""
    return "" if value is None else str(value).strip()


def as_number(value):
    """数字字符串 -> int/float；空串 -> None；其余原样返回。"""
    if isinstance(value, str):
        raw = value.strip()
        if not raw:
            return None
        try:
            number = float(raw)
        except ValueError:
            return value
        return int(number) if number.is_integer() else number
    return value


def positive_number(value):
    """正数（含数字字符串）-> int/float；其余（空、文本、0、负数）-> None。"""
    if isinstance(value, str):
        raw = value.strip()
        if not raw:
            return None
        try:
            number = float(raw)
        except ValueError:
            return None
    elif isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    else:
        number = float(value)
    if number <= 0:
        return None
    return int(number) if number.is_integer() else number


def num2(value):
    """保留 2 位小数，整数值写成整数（与参考表的写法一致）。"""
    rounded = round(float(value), 2)
    return int(rounded) if float(rounded).is_integer() else rounded


def norm_dims(raw):
    """'1050 x 1050 x 1150 mm' / '1050*1050*1150' -> '1050*1050*1150'。"""
    chunks = []
    for piece in re.split(r"[*xX×]", raw):
        digits = re.sub(r"[^0-9.]", "", piece)
        if digits:
            chunks.append(digits)
    return "*".join(chunks)


# ----------------------------------------------------------------------------
# 读取 + 校验唛头
# ----------------------------------------------------------------------------
def read_marks(path, errors):
    """读取唛头并校验，返回 [{pallet, sheet, parts, gross, dims}, ...]。"""
    wb = openpyxl.load_workbook(path, data_only=True)
    names = wb.sheetnames
    total = len(names)
    if total == 0:
        errors.append("唛头文件里没有任何 sheet")
        return []

    pallets = []
    lot_seen = {}
    for index, name in enumerate(names):
        ws = wb[name]
        pallet = index + 1
        where = f"唛头 sheet《{name}》"

        # 8A / 8B：LOT NUMBER = 托盘号-托盘总数
        label = text(ws.cell(row=LOT_ROW, column=1).value)
        if label != LOT_LABEL:
            errors.append(f"{where}: A8 应为 {LOT_LABEL!r}，实际为 {label!r}")
        lot_raw = text(ws.cell(row=LOT_ROW, column=2).value)
        matched = LOT_RE.match(lot_raw)
        if not matched:
            errors.append(
                f"{where}: B8 的 LOT NUMBER {lot_raw!r} 不符合「托盘号-托盘总数」格式"
            )
        else:
            lot_no, lot_total = int(matched.group(1)), int(matched.group(2))
            if lot_no != pallet:
                errors.append(
                    f"{where}: LOT NUMBER 托盘号 {lot_no} 与该 sheet 的顺序 {pallet} 不一致"
                )
            if lot_total != total:
                errors.append(
                    f"{where}: LOT NUMBER 托盘总数 {lot_total} 与 sheet 总数 {total} 不一致"
                )
            lot_seen.setdefault(lot_no, []).append(name)

        # 定位 WEIGHT (IN KG) 行：它之前是 part 行，之后是托盘的其它信息
        rows = list(ws.iter_rows(min_row=FIRST_PART_ROW, max_row=max(ws.max_row, FIRST_PART_ROW)))
        weight_row = next((r for r in rows if text(r[0].value) == WEIGHT_LABEL), None)
        if weight_row is None:
            errors.append(f"{where}: 未找到 {WEIGHT_LABEL!r} 行，无法确定 part 行范围")
            part_rows = []
            gross = None
        else:
            part_rows = [r for r in rows if r[0].row < weight_row[0].row]
            gross = positive_number(weight_row[1].value)
        dims_row = next((r for r in rows if text(r[0].value) == DIMS_LABEL), None)
        dims = text(dims_row[1].value) if dims_row is not None else ""

        parts = []
        for row in part_rows:
            part_number = text(row[0].value)
            if not part_number:
                continue
            qty = positive_number(row[2].value)
            if qty is None:
                errors.append(
                    f"{where} 第{row[0].row}行: part {part_number} 的数量 {row[2].value!r} 不是正数"
                )
                continue
            parts.append((part_number, qty))

        if not parts:
            errors.append(f"{where}: 没有解析到任何 part 数据")
        duplicated = [pn for pn, count in Counter(pn for pn, _ in parts).items() if count > 1]
        if duplicated:
            errors.append(f"{where}: part-number 重复: {', '.join(duplicated)}")
        if gross is None:
            errors.append(f"{where}: {WEIGHT_LABEL} 的总重不是正数")
        if not dims:
            errors.append(f"{where}: 未找到 {DIMS_LABEL!r} 尺寸")

        pallets.append({
            "pallet": pallet,
            "sheet": name,
            "parts": parts,
            "gross": gross,
            "dims": norm_dims(dims),
        })

    for lot_no, sheets in lot_seen.items():
        if len(sheets) > 1:
            errors.append(f"LOT NUMBER 托盘号 {lot_no} 重复出现在 sheet: {', '.join(sheets)}")
    return pallets


# ----------------------------------------------------------------------------
# 读取 + 校验 ASN 数据 sheet / 单位重量 sheet
# ----------------------------------------------------------------------------
def read_asn_parts(path, sheet_name, errors):
    """读取 ASN 数据 sheet，返回 {part_number: [31 个单元格值]}。"""
    wb = openpyxl.load_workbook(path, data_only=True)
    if sheet_name not in wb.sheetnames:
        errors.append(f"ASN 文件里没有名为《{sheet_name}》的 sheet")
        return {}
    ws = wb[sheet_name]

    actual = [text(cell.value) for cell in ws[1]]
    while actual and actual[-1] == "":
        actual.pop()
    if actual != HEADERS:
        if len(actual) != len(HEADERS):
            errors.append(
                f"ASN《{sheet_name}》列数不符: 期望 {len(HEADERS)} 列，实际 {len(actual)} 列"
            )
        for i, (got, want) in enumerate(zip(actual, HEADERS)):
            if got != want:
                letter = get_column_letter(i + 1)
                errors.append(
                    f"ASN《{sheet_name}》{letter}1 列名不符: 期望 {want!r}，实际 {got!r}"
                )

    parts = {}
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        values = [cell.value for cell in row]
        values += [None] * (len(HEADERS) - len(values))
        values = values[:len(HEADERS)]
        if all(text(value) == "" for value in values):
            continue
        part_number = text(values[COL["part_number"]])
        if not part_number:
            errors.append(f"ASN《{sheet_name}》第{row[0].row}行: part_number 为空")
            continue
        if part_number in parts:
            errors.append(f"ASN《{sheet_name}》第{row[0].row}行: part_number {part_number} 重复")
            continue
        parts[part_number] = values

    if not parts:
        errors.append(f"ASN《{sheet_name}》: 没有解析到任何 part-number")
    return parts


def read_unit_weights(path, sheet_name, errors):
    """读取单位重量 sheet，返回 {part_number: 单位重量(float)}。"""
    wb = openpyxl.load_workbook(path, data_only=True)
    if sheet_name not in wb.sheetnames:
        errors.append(f"ASN 文件里没有名为《{sheet_name}》的 sheet")
        return {}
    ws = wb[sheet_name]

    extra = [c for c in ws.iter_cols(min_col=3) if any(text(cell.value) for cell in c)]
    if extra:
        errors.append(
            f"单位重量《{sheet_name}》: 应只有 2 列数据，实际存在第 3 列及以后的数据"
        )

    weights = {}
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        raw_name, raw_weight = row[0].value, row[1].value
        if text(raw_name) == "" and text(raw_weight) == "":
            continue
        part_number = text(raw_name)
        if not part_number:
            errors.append(f"单位重量《{sheet_name}》第{row[0].row}行: A 列 part-number 为空")
            continue
        if not isinstance(raw_weight, (int, float)) or isinstance(raw_weight, bool) or raw_weight <= 0:
            errors.append(
                f"单位重量《{sheet_name}》第{row[0].row}行: {part_number} 的单位重量 "
                f"{raw_weight!r} 不是正浮点数"
            )
            continue
        if part_number in weights:
            errors.append(f"单位重量《{sheet_name}》第{row[0].row}行: part-number {part_number} 重复")
            continue
        weights[part_number] = float(raw_weight)

    if not weights:
        errors.append(f"单位重量《{sheet_name}》: 没有解析到任何 part-number")
    return weights


# ----------------------------------------------------------------------------
# 组装输出行
# ----------------------------------------------------------------------------
def build_rows(pallets, asn_parts, weights, errors):
    """把唛头展开成 ASN_Template 的数据行（每行 31 个值）。"""
    rows = []
    for pallet in pallets:
        pallet_id = pallet["pallet"]
        net_total = 0.0
        for part_number, qty in pallet["parts"]:
            if part_number in weights:
                net_total += qty * weights[part_number]

        first_of_pallet = True
        for part_number, qty in pallet["parts"]:
            base = asn_parts.get(part_number)
            if base is None:
                errors.append(
                    f"托盘 {pallet_id}（sheet《{pallet['sheet']}》）的 part-number "
                    f"{part_number} 在 ASN 数据 sheet 中不存在"
                )
                continue
            if part_number not in weights:
                errors.append(
                    f"托盘 {pallet_id} 的 part-number {part_number} 在单位重量 sheet 中不存在"
                )
                continue

            row = []
            for i, name in enumerate(HEADERS):
                value = base[i]
                if name in NUMERIC_COLS:
                    row.append(as_number(value))
                else:
                    row.append(None if text(value) == "" else value)

            row[COL["items_per_package"]] = qty
            row[COL["dimensions"]] = pallet["dims"]
            row[COL["package_id"]] = pallet_id
            if first_of_pallet:
                # 净重/毛重只写在每个托盘的首行
                row[COL["net_weight_per_package"]] = num2(net_total)
                row[COL["gross_weight_per_package"]] = num2(pallet["gross"])
                first_of_pallet = False
            else:
                row[COL["net_weight_per_package"]] = None
                row[COL["gross_weight_per_package"]] = None
            rows.append(row)
    return rows


def write_output(path, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = OUT_SHEET
    ws.append(HEADERS)
    for row in rows:
        ws.append(row)
    wb.save(path)


# ----------------------------------------------------------------------------
# 入口
# ----------------------------------------------------------------------------
def main():
    parser = argparse.ArgumentParser(
        description="根据唛头 + ASN 导出生成可上传 ASN 系统的 ASN_Template 表格"
    )
    parser.add_argument("--mark-path", help="唛头 Excel 路径（每个 sheet = 1 个托盘）")
    parser.add_argument("--asn-path", help="ASN 导出的 Excel 路径")
    parser.add_argument("--asn-sheet", required=True, metavar="NAME",
                        help="ASN 数据 sheet 名，如 \"ASN_Template (2)\"")
    parser.add_argument("--weight-sheet", required=True, metavar="NAME",
                        help="单位重量 sheet 名，如 \"Sheet1\"（无表头 2 列）")
    parser.add_argument("--out", required=True, metavar="PATH",
                        help="输出 Excel 路径（必须与输入路径不同）")
    args = parser.parse_args()

    marks_path = os.path.abspath(args.mark_path)
    asn_path = os.path.abspath(args.asn_path)
    out_path = os.path.abspath(args.out)

    for label, path in (("唛头", marks_path), ("ASN", asn_path)):
        if not os.path.isfile(path):
            sys.exit(f"{label}文件不存在: {path}")
    if os.path.realpath(out_path) in (os.path.realpath(marks_path), os.path.realpath(asn_path)):
        sys.exit(f"输出路径不能与输入路径相同: {out_path}")

    errors = []
    pallets = read_marks(marks_path, errors)
    asn_parts = read_asn_parts(asn_path, args.asn_sheet, errors)
    weights = read_unit_weights(asn_path, args.weight_sheet, errors)

    # 交叉校验：ASN 中每个 part-number 都要能在单位重量表中查到
    for part_number in asn_parts:
        if part_number not in weights:
            errors.append(f"ASN 的 part-number {part_number} 在单位重量 sheet 中不存在")

    rows = []
    if pallets and asn_parts and weights:
        rows = build_rows(pallets, asn_parts, weights, errors)

    if errors:
        print(f"校验未通过，共 {len(errors)} 个问题：", file=sys.stderr)
        for i, message in enumerate(errors, 1):
            print(f"  {i}. {message}", file=sys.stderr)
        sys.exit(1)

    write_output(out_path, rows)
    print(f"已生成 {out_path}")
    print(f"  托盘数: {len(pallets)}，数据行数: {len(rows)}")


if __name__ == "__main__":
    main()
