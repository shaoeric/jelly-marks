#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
N20 唛头生成器（Excel 输出，一托一 sheet）
=========================================

输入
----
1. 线束计算表（如 3762062_SH.xlsx）的 sheet "Combined POS"
   - 表头行: B 列 = "Combiner Box Type"，C 列 = "Home Run Harness Lengths"
     （C..S 共 17 个"位置"列，下一行为 1..17），其后一列 = "SUM"
   - 表头行右侧: 形如 T-N20-JG-0125 的工艺卡号列，其**下方一行的数字**就是该
     卡号对应的线束长度（英尺）：0125=24 / 0128=48 / 0129=72 / 0130=96 /
     0131=120 / 0132=144
   - 数据区每行 = 一个 Combiner 串（CMB-xx-yy），给出该串各长度的用量
2. 装箱明细（如 Pallet detail-*.xlsx）的 sheet "Sheet1 (2)"
   - 表头行: A 列 = "Block No."，其后字段名与顺序见 PALLET_HEADER
   - 数据区每个"有 No. 的行"= 一个托盘，其下"No. 为空"的行属于同一个托盘

输出
----
一个 sheet = 一个托盘（sheet 名 = 托盘号），字段：
抬头(整票固定信息) / LOT NUMBER "托盘号-总托数" / 明细行(PART NUMBER、PART
DESCRIPTION、SH LENGTH、QUANTITY/PN) / WEIGHT / DIMENSIONS / TYPE OF STORAGE /
STACKABILITY / BLOCK NUMBER / PART PHOTO

计算规则
--------
- SH LENGTH : 工艺卡号在长度表里 -> "24 FT"…；否则（PPH、JUMPER）填 "/"
- 明细行    : 每个托盘明细行一行；PN/DESC 为空的续行（如 48FT 线束行）沿用上一行
- WEIGHT    : 该托盘各行 Weight(kg) 求和（该列即所在托盘的毛重）；仅当整托都
              没有 Weight 时才按公式补算：数量 x Unit NW + 2kg x 箱数
              (+ 55kg 托盘自重，只计一次，记在该托盘第一行有"Pallet size"的行上)
- DIMENSIONS: 该托盘第一个非空的 Pallet size(mm) + " mm"
- BLOCK     : "INV xx"（A 列分段值）；若该托盘含 SH 线束，再追加这些线束服务的
              CMB 串（按 Combined POS 中 CMB 顺序、按长度分别先进先出累计分配）

用法
----
    python3 generate_n20.py \
        --combined "3762062_SH.xlsx" --combined-sheet "Combined POS" \
        --pallet   "Pallet detail-N20-260824-1st.xlsx" --pallet-sheet "Sheet1 (2)" \
        --out "唛头 -N20-260824-1st.xlsx" [--logo logo.png]

依赖: pip install openpyxl
"""

import argparse
import os
import sys
from collections import OrderedDict

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("缺少 openpyxl，请先执行: pip install openpyxl")


# ----------------------------------------------------------------------------
# 配置：整票固定信息（进货/出货资料，两个输入表里都没有，按需修改或用参数覆盖）
# ----------------------------------------------------------------------------
CONFIG = {
    "supplier_id": "4047270",
    "country_of_origin": "THAILAND",
    "project_number": "4096638",
    "project_name": "Hinds MS - 210MW - eBOS CBS",
    "so_number": "42194",
    "po_number": "UAE93474",
    "type_of_storage": "INDOOR",
    "stackability": "NO",
    "carton_weight_kg": 2.0,          # 纸箱自重（Weight 缺失时补算用）
    "pallet_weight_kg": 55.0,         # 托盘自重（同上）
}

# 装箱明细表头（字段名 + 顺序，逐列校验；A..S 必填）
PALLET_HEADER = [
    "Block No.", "No.", "物料编码", "工艺卡号", "Item", "PN", "DESC",
    "Drawing No.", "QTY(set)", "Unit NW(kg)", "QTY/Carton", "Carton QTY",
    "Carton/PLT", "PLT QTY", "Rest carton", "Tail carton QTY", "Carton size(mm)",
    "Pallet size(mm)", "Weight(kg)",
]
# 表尾附加列（参考文件里第一段表头没有这一列，故有则校验、无则跳过）
PALLET_HEADER_OPTIONAL = ["Tail pallet"]
# 线束计算表表头
COMBINED_BOX_COL = "Combiner Box Type"
COMBINED_LEN_COL = "Home Run Harness Lengths"
COMBINED_SUM_COL = "SUM"
HARNESS_PREFIX = "T-N20-JG-"
COMBINED_POSITIONS = 17                 # Home Run Harness Lengths 的位置数
CMB_PER_LINE = 8                        # BLOCK 里 CMB 列表每行放几个

# 装箱明细列（1 基）
C_BLOCK, C_NO, C_CARD, C_PN, C_DESC = 1, 2, 4, 6, 7
C_QTY, C_UNIT, C_CARTONS = 9, 10, 12
C_PALLET_SIZE, C_WEIGHT = 18, 19
# 线束计算表列（1 基）
K_CARD, K_BOX = 2, 3

FONT_NAME = "Neue Montreal"             # 与参考唛头一致；本机没有则用默认字体
MARK_OUT = "唛头.xlsx"


# ----------------------------------------------------------------------------
# 通用工具
# ----------------------------------------------------------------------------
def resolve_input(path):
    """返回输入文件的绝对路径；不存在则退出。"""
    if not isinstance(path, str) or not path.strip():
        sys.exit("输入文件路径为空")
    raw = os.path.expanduser(path.strip())
    for seg in raw.replace("\\", "/").split("/"):
        if seg == "..":
            sys.exit("非法路径（含 ..）: %s" % raw)
    p = os.path.normpath(os.path.abspath(raw))
    if not os.path.isfile(p):
        sys.exit("找不到输入文件: %s" % p)
    return p


def get_sheet(wb, name, path):
    """按名字取 sheet；不存在则列出可选 sheet 名后退出。"""
    if name not in wb.sheetnames:
        sys.exit("文件 %s 里没有 sheet %r，可选: %s"
                 % (os.path.basename(path), name, wb.sheetnames))
    return wb[name]


def norm_text(value):
    """表头文本归一化（合并空白），非字符串返回 None。"""
    if not isinstance(value, str):
        return None
    return " ".join(value.replace("\n", " ").split())


def num(value):
    """数值化；空值/斜杠返回 None。"""
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip().replace(",", "")
    if not s or s == "/":
        return None
    try:
        return float(s)
    except ValueError:
        return None


def fmt_num(x):
    """去掉多余小数位: 386.15 / 389 / 552.01"""
    s = ("%.2f" % round(float(x), 2)).rstrip("0").rstrip(".")
    return s if s else "0"


def fmt_qty(q):
    """数量：整数去掉小数点"""
    f = float(q)
    return str(int(f)) if abs(f - int(f)) < 1e-9 else fmt_num(f)


# ----------------------------------------------------------------------------
# 校验 1：装箱明细字段名 + 字段顺序
# ----------------------------------------------------------------------------
def validate_pallet_sheet(ws):
    """校验装箱明细表头（字段名与顺序），返回所有表头行号。"""
    header_rows = [r for r in range(1, ws.max_row + 1)
                   if norm_text(ws.cell(r, C_BLOCK).value) == PALLET_HEADER[0]]
    if not header_rows:
        sys.exit("装箱明细校验失败：找不到表头行（A 列须为 %r）" % PALLET_HEADER[0])
    for r in header_rows:
        for i, want in enumerate(PALLET_HEADER, start=1):
            got = norm_text(ws.cell(r, i).value)
            if got != want:
                sys.exit("装箱明细校验失败：%s%d 字段应为 %r，实际为 %r"
                         % (get_column_letter(i), r, want, got))
        for j, want in enumerate(PALLET_HEADER_OPTIONAL,
                                 start=len(PALLET_HEADER) + 1):
            got = norm_text(ws.cell(r, j).value)
            if got is not None and got != want:
                sys.exit("装箱明细校验失败：%s%d 字段应为 %r 或留空，实际为 %r"
                         % (get_column_letter(j), r, want, got))
    return header_rows


# ----------------------------------------------------------------------------
# 校验 2：线束计算表字段名 + 字段顺序，并取出 工艺卡号 -> 长度
# ----------------------------------------------------------------------------
def validate_combined_sheet(ws):
    """校验线束计算表表头，返回 (表头行号, {列号: (工艺卡号, 长度FT)})。"""
    hdr = None
    for r in range(1, min(ws.max_row, 50) + 1):
        if norm_text(ws.cell(r, K_CARD).value) == COMBINED_BOX_COL:
            hdr = r
            break
    if hdr is None:
        sys.exit("线束计算表校验失败：找不到表头行（B 列须为 %r）" % COMBINED_BOX_COL)

    def expect(row, col, want):
        got = norm_text(ws.cell(row, col).value)
        if got != want:
            sys.exit("线束计算表校验失败：%s%d 应为 %r，实际 %r"
                     % (get_column_letter(col), row, want, got))

    expect(hdr, K_BOX, COMBINED_LEN_COL)
    for i in range(COMBINED_POSITIONS):                       # 位置 1..17 顺序
        col = K_BOX + i
        got = ws.cell(hdr + 1, col).value
        if got != i + 1:
            sys.exit("线束计算表校验失败：%s%d 位置编号应为 %d，实际 %r"
                     % (get_column_letter(col), hdr + 1, i + 1, got))
    expect(hdr, K_BOX + COMBINED_POSITIONS, COMBINED_SUM_COL)

    harness = OrderedDict()                                   # 列号 -> (卡号, 长度)
    lengths = []
    for col in range(K_BOX + 1 + COMBINED_POSITIONS, ws.max_column + 1):
        card = norm_text(ws.cell(hdr, col).value)
        if not (card and card.startswith(HARNESS_PREFIX)):
            continue
        ft = num(ws.cell(hdr + 1, col).value)
        if ft is None:
            sys.exit("线束计算表校验失败：%s%d(%s) 下方缺少线束长度(英尺)"
                     % (get_column_letter(col), hdr + 1, card))
        if card in [c for c, _ in harness.values()]:
            sys.exit("线束计算表校验失败：工艺卡号 %s 重复" % card)
        if any(ft <= x for x in lengths):
            sys.exit("线束计算表校验失败：线束长度需按从小到大排列，%s 处为 %s FT"
                     % (card, fmt_num(ft)))
        lengths.append(ft)
        harness[col] = (card, ft)
    if not harness:
        sys.exit("线束计算表校验失败：找不到任何 %s* 线束列" % HARNESS_PREFIX)
    return hdr, harness


# ----------------------------------------------------------------------------
# 读取线束计算表：每个 CMB 串的各长度用量
# ----------------------------------------------------------------------------
def read_cmb_demand(ws, hdr, harness):
    """返回 (CMB 顺序, {长度FT: [(CMB, 数量), ...]})。"""
    order = []
    demand = OrderedDict((ft, []) for _, ft in harness.values())
    for r in range(hdr + 1, ws.max_row + 1):
        cmb = norm_text(ws.cell(r, K_CARD).value)
        if not cmb:
            continue
        order.append(cmb)
        for col, (_, ft) in harness.items():
            qty = num(ws.cell(r, col).value)
            if qty:
                demand[ft].append((cmb, qty))
    if not order:
        sys.exit("线束计算表校验失败：表头下方没有 CMB 数据行")
    return order, demand


# ----------------------------------------------------------------------------
# 读取装箱明细：按托盘分组
# ----------------------------------------------------------------------------
def read_pallets(ws, header_rows, length_of):
    """把装箱明细按托盘分组；length_of: 工艺卡号 -> 长度FT/None。"""
    stops = list(header_rows[1:]) + [ws.max_row + 1]
    pallets = OrderedDict()
    for hdr, stop in zip(header_rows, stops):
        inv, cur, seen_size = None, None, False
        last_pn, last_desc = None, None
        for r in range(hdr + 1, stop):
            block = norm_text(ws.cell(r, C_BLOCK).value)
            if block:
                inv = block
            qty = num(ws.cell(r, C_QTY).value)
            if qty is None:
                continue
            no = num(ws.cell(r, C_NO).value)
            if no is not None:
                no = int(no)
                if no in pallets:
                    sys.exit("装箱明细校验失败：托盘号 %d 重复出现（第 %d 行）；"
                             "请确认 --pallet-sheet 是按托盘拆分的表（如 Sheet1 (2)）"
                             % (no, r))
                pallets[no] = {"no": no, "inv": inv, "items": [],
                               "weight": 0.0, "calc": 0.0, "has_weight": False,
                               "dims": ""}
                cur, seen_size = no, False
            if cur is None:
                sys.exit("装箱明细校验失败：第 %d 行有数量但没有托盘号(No.)" % r)

            card = norm_text(ws.cell(r, C_CARD).value)
            p = pallets[cur]
            pn = norm_text(ws.cell(r, C_PN).value)
            desc = norm_text(ws.cell(r, C_DESC).value)
            if pn is None:                       # 续行（如 48FT 线束）沿用上一行 PN/DESC
                pn, desc = last_pn, last_desc
            else:
                last_pn, last_desc = pn, desc
            if pn is None:
                sys.exit("装箱明细校验失败：第 %d 行缺少 PN，且上一行没有可沿用的 PN" % r)

            ft = length_of.get(card)
            size = norm_text(ws.cell(r, C_PALLET_SIZE).value)
            first_of_size = bool(size) and not seen_size
            unit = num(ws.cell(r, C_UNIT).value) or 0.0
            cartons = num(ws.cell(r, C_CARTONS).value) or 0.0
            # 输入已有毛重就照用；整托都没有毛重时才按公式补算
            calc = qty * unit + CONFIG["carton_weight_kg"] * cartons
            if first_of_size:
                calc += CONFIG["pallet_weight_kg"]
            if size:
                seen_size = True
                if not p["dims"]:
                    p["dims"] = size
            w = num(ws.cell(r, C_WEIGHT).value)
            p["items"].append({"pn": pn, "desc": desc or "", "card": card,
                               "ft": ft, "qty": qty, "row": r, "w": w, "calc": calc})
            p["weight"] += w if w is not None else 0.0
            p["calc"] += calc
            p["has_weight"] = p["has_weight"] or (w is not None)
    for p in pallets.values():
        if p["inv"] is None:
            sys.exit("装箱明细校验失败：托盘 %s 找不到 Block No." % p["no"])
        if not p["items"]:
            sys.exit("装箱明细校验失败：托盘 %d 没有明细行" % p["no"])
        if not p["dims"]:
            sys.exit("装箱明细校验失败：托盘 %d 没有 Pallet size(mm)" % p["no"])
        p["weight"] = round(p["weight"] if p["has_weight"] else p["calc"], 2)
    return pallets


# ----------------------------------------------------------------------------
# CMB 归属：按长度分别先进先出累计
# ----------------------------------------------------------------------------
def assign_cmb_blocks(pallets, demand, warnings):
    """给每个含 SH 线束的托盘标出它服务到的 CMB 串。"""
    ptr = {ft: 0 for ft in demand}
    rem = {ft: 0.0 for ft in demand}
    for p in pallets.values():
        touched = OrderedDict()
        for it in p["items"]:
            ft = it["ft"]
            if ft is None:
                continue
            need = it["qty"]
            while need > 0 and ptr[ft] < len(demand[ft]):
                cmb, n = demand[ft][ptr[ft]]
                if rem[ft] == 0:
                    rem[ft] = n
                take = min(need, rem[ft])
                if take > 0:
                    touched[cmb] = None
                need -= take
                rem[ft] -= take
                if rem[ft] <= 0:
                    ptr[ft] += 1
            if need > 0:
                warnings.append("托盘 %s 的 %s FT 线束有 %s 支超出了 Combined POS 的总需求量"
                                % (p["no"], fmt_num(ft), fmt_qty(need)))
        p["cmbs"] = list(touched)
    return pallets


def block_text(inv, cmbs):
    """BLOCK NUMBER 文本：INV 段 + CMB 串（每行 8 个）。"""
    if not cmbs:
        return inv
    lines = ["，".join(cmbs[i:i + CMB_PER_LINE])
             for i in range(0, len(cmbs), CMB_PER_LINE)]
    return "%s，%s" % (inv, "\n".join(lines))


# ----------------------------------------------------------------------------
# 写出唛头 sheet
# ----------------------------------------------------------------------------
def write_mark_sheet(ws, cfg, lot, total, items, weight, dims, block, logo=None):
    """一个托盘 = 一个 sheet（版式对齐参考唛头）。"""
    font = Font(name=FONT_NAME, size=10)
    font_bold = Font(name=FONT_NAME, size=10, bold=True)
    med = Side(style="medium")
    grid = Border(left=med, right=med, top=med, bottom=med)
    for col, w in (("A", 24.5), ("B", 46.4), ("C", 15.56), ("D", 23.5)):
        ws.column_dimensions[col].width = w
    ws.sheet_view.showGridLines = False

    # 第 1 行：LOGO 区
    ws.row_dimensions[1].height = 40
    ws.merge_cells("B1:D1")
    if logo:
        # 只有需要贴 LOGO 时才用 openpyxl 的图片功能（要 Pillow）
        from openpyxl.drawing.image import Image
        ws.add_image(Image(logo), "A1")

    # 第 2-8 行：抬头
    ws.merge_cells("A2:D2")
    ws["A2"] = "SUPPLIER ID:   %s" % cfg["supplier_id"]
    ws["A2"].font = font_bold
    ws.merge_cells("A3:D3")
    ws["A3"] = "COUNTRY OF ORIGIN: %s" % cfg["country_of_origin"]
    ws["A3"].font = font
    head = [
        ("PROJECT NUMBER", cfg["project_number"]),
        ("PROJECT NAME", cfg["project_name"]),
        ("SO NUMBER", cfg["so_number"]),
        ("PO NUMBER", cfg["po_number"]),
        ("LOT NUMBER", "%d-%d" % (lot, total)),
    ]
    r = 4
    for label, value in head:
        ws.cell(r, 1, label).font = font
        ws.cell(r, 2, value).font = font_bold if label == "LOT NUMBER" else font
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        ws.row_dimensions[r].height = 25
        r += 1

    # 第 9 行：明细表头
    for col, text in ((1, "PART NUMBER"), (2, "PART DESCRIPTION"),
                      (3, "SH LENGTH"), (4, "QUANTITY/PN")):
        ws.cell(r, col, text).font = font
    ws.row_dimensions[r].height = 25
    r += 1

    # 明细行
    for it in items:
        ws.cell(r, 1, it["pn"]).font = font
        ws.cell(r, 2, it["desc"]).font = font
        ws.cell(r, 3, "%s FT" % fmt_num(it["ft"]) if it["ft"] else "/").font = font
        ws.cell(r, 4, fmt_qty(it["qty"])).font = font
        ws.row_dimensions[r].height = 25
        r += 1

    # 尾部字段
    tail = [
        ("WEIGHT (IN KG)", fmt_num(weight)),
        ("DIMENSIONS (MM) LxWxH", "%s mm" % dims),
        ("TYPE OF STORAGE", cfg["type_of_storage"]),
        ("STACKABILITY", cfg["stackability"]),
        ("BLOCK NUMBER", block),
    ]
    for label, value in tail:
        ws.cell(r, 1, label).font = font
        cell = ws.cell(r, 2, value)
        cell.font = font_bold if label == "WEIGHT (IN KG)" else font
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        ws.row_dimensions[r].height = max(25, 15 * (str(value).count("\n") + 1) + 10)
        r += 1

    # PART PHOTO 留白
    ws.cell(r, 1, "PART PHOTO").font = font
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    ws.row_dimensions[r].height = 146
    last = r

    for row in ws.iter_rows(min_row=1, max_row=last, min_col=1, max_col=4):
        for c in row:
            c.border = grid
            if c.alignment.vertical is None:
                c.alignment = Alignment(vertical="center")
    ws.print_area = "A1:D%d" % last
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = 9                      # A4
    return last


def make_workbook(path, pallets, cfg, logo=None):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    total = len(pallets)
    for p in pallets.values():
        ws = wb.create_sheet(title=str(p["no"]))
        write_mark_sheet(ws, cfg, p["no"], total, p["items"], p["weight"],
                         p["dims"], block_text(p["inv"], p["cmbs"]), logo)
    wb.save(path)
    return total


# ----------------------------------------------------------------------------
# 主流程
# ----------------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser(description="N20 唛头生成器（一托一 sheet）")
    ap.add_argument("--combined", required=True, help="线束计算表 xlsx（如 3762062_SH.xlsx）")
    ap.add_argument("--combined-sheet", default="Combined POS", help="线束计算表 sheet 名")
    ap.add_argument("--pallet", required=True, help="装箱明细 xlsx（如 Pallet detail-*.xlsx）")
    ap.add_argument("--pallet-sheet", default="Sheet1 (2)", help="装箱明细 sheet 名")
    ap.add_argument("--out", default=MARK_OUT, help="输出唛头文件名（默认 %s）" % MARK_OUT)
    ap.add_argument("--logo", default=None, help="可选的 LOGO 图片，放在每个 sheet 第 1 行")
    ap.add_argument("--supplier-id", default=CONFIG["supplier_id"])
    ap.add_argument("--country-of-origin", default=CONFIG["country_of_origin"])
    ap.add_argument("--project-number", default=CONFIG["project_number"])
    ap.add_argument("--project-name", default=CONFIG["project_name"])
    ap.add_argument("--so-number", default=CONFIG["so_number"])
    ap.add_argument("--po-number", default=CONFIG["po_number"])
    args = ap.parse_args()

    cfg = {
        "supplier_id": args.supplier_id,
        "country_of_origin": args.country_of_origin,
        "project_number": args.project_number,
        "project_name": args.project_name,
        "so_number": args.so_number,
        "po_number": args.po_number,
        "type_of_storage": CONFIG["type_of_storage"],
        "stackability": CONFIG["stackability"],
    }

    combined_path = resolve_input(args.combined)
    pallet_path = resolve_input(args.pallet)
    logo = resolve_input(args.logo) if args.logo else None

    # 1) 线束计算表：校验表头 -> 工艺卡号 -> 长度 / 各 CMB 需求
    wb = openpyxl.load_workbook(combined_path, data_only=True)
    cws = get_sheet(wb, args.combined_sheet, combined_path)
    hdr, harness = validate_combined_sheet(cws)
    _, demand = read_cmb_demand(cws, hdr, harness)
    length_of = {card: ft for card, ft in harness.values()}
    print("线束计算表: %s[%s] 表头第 %d 行，线束卡号 %s"
          % (os.path.basename(combined_path), args.combined_sheet, hdr,
             {c: "%s FT" % fmt_num(f) for c, f in length_of.items()}))

    # 2) 装箱明细：校验表头 -> 按托盘分组
    wb2 = openpyxl.load_workbook(pallet_path, data_only=True)
    pws = get_sheet(wb2, args.pallet_sheet, pallet_path)
    header_rows = validate_pallet_sheet(pws)
    pallets = read_pallets(pws, header_rows, length_of)
    print("装箱明细: %s[%s] 表头 %d 段，托盘 %s 个（%d~%d）"
          % (os.path.basename(pallet_path), args.pallet_sheet, len(header_rows),
             len(pallets), min(pallets), max(pallets)))

    # 3) CMB 归属 -> 4) 写唛头
    warnings = []
    assign_cmb_blocks(pallets, demand, warnings)
    total = make_workbook(args.out, pallets, cfg, logo)

    for p in pallets.values():
        print("托盘 %-3s | LOT %s-%s | %s | %s kg | %s"
              % (p["no"], p["no"], total, p["dims"],
                 fmt_num(p["weight"]), block_text(p["inv"], p["cmbs"]).replace("\n", "")))
    for w in warnings:
        print("warning: %s" % w)
    print("输出唛头: %s（%d 个 sheet）" % (args.out, total))


if __name__ == "__main__":
    main()
